from __future__ import annotations

import csv
import hashlib
import hmac
import json
import os
import subprocess
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from binance.client import Client
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
HISTORY_FILE = BASE_DIR / "portfolio_history.csv"
WORKBOOK_FILE = BASE_DIR / "PlanilhaCalculaCripto.xlsx"
KNOWN_TRADE_SYMBOLS_FILE = BASE_DIR / "trade_symbols.json"
EXCEL_REFRESH_SCRIPT = BASE_DIR / "refresh_excel.ps1"
BINANCE_API_URL = "https://api.binance.com"
CACHE_SECONDS = 60
TRADE_QUOTES = {"USDT", "USDC", "FDUSD", "BRL"}
USD_LIKE_QUOTES = {"USDT", "USDC", "FDUSD"}

load_dotenv(BASE_DIR / ".env")
load_dotenv(BASE_DIR.parent / ".env")

api_key = os.getenv("BINANCE_API_KEY")
api_secret = os.getenv("BINANCE_API_SECRET")
client = Client(api_key, api_secret) if api_key and api_secret else None

app = FastAPI(title="Crypto Investments Dashboard", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:8000", "http://127.0.0.1:8000"],
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

_cache_lock = threading.Lock()
_refresh_lock = threading.Lock()
_dashboard_cache: dict[str, Any] = {"data": None, "updated_monotonic": 0.0}


def require_client() -> Client:
    if client is None:
        raise RuntimeError(
            "Credenciais da Binance não encontradas. Configure BINANCE_API_KEY e "
            "BINANCE_API_SECRET no arquivo .env da pasta Crypto."
        )
    return client


def normalize_asset(asset: str) -> str:
    return asset[2:] if asset.startswith("LD") else asset


def normalize_for_lookup(asset: str) -> str:
    return "SHIB" if asset == "SHIB2" else asset


def get_current_price(symbol: str) -> float | None:
    try:
        return float(require_client().get_symbol_ticker(symbol=symbol)["price"])
    except Exception:
        return None


def get_price_in_usdt(asset: str) -> float | None:
    if asset in {"USDT", "FDUSD", "USDC"}:
        return 1.0
    if asset == "BRL":
        usdt_brl = get_current_price("USDTBRL")
        return 1 / usdt_brl if usdt_brl else None
    return get_current_price(f"{normalize_for_lookup(asset)}USDT")


def signed_request(url_path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    require_client()
    assert api_key and api_secret
    params = dict(payload or {})
    params["timestamp"] = int(time.time() * 1000)
    params["recvWindow"] = 10_000
    query_string = "&".join(f"{key}={value}" for key, value in params.items())
    params["signature"] = hmac.new(
        api_secret.encode("utf-8"), query_string.encode("utf-8"), hashlib.sha256
    ).hexdigest()
    response = requests.get(
        BINANCE_API_URL + url_path,
        headers={"X-MBX-APIKEY": api_key},
        params=params,
        timeout=30,
    )
    response.raise_for_status()
    payload_json = response.json()
    if isinstance(payload_json, dict) and payload_json.get("code", 0) < 0:
        raise RuntimeError(payload_json.get("msg", "Erro retornado pela Binance"))
    return payload_json


def read_known_trade_symbols() -> set[str]:
    """Keep querying pairs that appeared before, even after an asset was sold."""
    symbols: set[str] = set()
    if KNOWN_TRADE_SYMBOLS_FILE.exists():
        try:
            symbols.update(
                str(symbol).upper()
                for symbol in json.loads(KNOWN_TRADE_SYMBOLS_FILE.read_text("utf-8"))
            )
        except Exception:
            pass
    if not WORKBOOK_FILE.exists():
        return symbols
    try:
        workbook = load_workbook(WORKBOOK_FILE, read_only=True, data_only=True)
        if "transactions" not in workbook.sheetnames:
            workbook.close()
            return symbols
        sheet = workbook["transactions"]
        symbols.update(
            str(row[1]).upper()
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if len(row) > 1 and row[1]
        )
        workbook.close()
        return symbols
    except Exception:
        return symbols


def read_cached_trades_from_workbook() -> list[dict[str, Any]]:
    """Read the last successful trade snapshot so API hiccups cannot erase it."""
    if not WORKBOOK_FILE.exists():
        return []
    try:
        workbook = load_workbook(WORKBOOK_FILE, read_only=True, data_only=True)
        if "transactions" not in workbook.sheetnames:
            workbook.close()
            return []
        rows = list(workbook["transactions"].iter_rows(values_only=True))
        workbook.close()
        if not rows:
            return []
        headers = [str(value or "").removeprefix("Column1.") for value in rows[0]]
        cached: list[dict[str, Any]] = []
        for values in rows[1:]:
            record = dict(zip(headers, values))
            if not record.get("symbol") or record.get("time") in (None, ""):
                continue
            symbol = str(record["symbol"]).upper()
            quote_asset = next(
                (quote for quote in TRADE_QUOTES if symbol.endswith(quote)), "USDT"
            )
            cached.append(
                {
                    "asset": str(record.get("asset") or asset_from_symbol(symbol)),
                    "symbol": symbol,
                    "quoteAsset": quote_asset,
                    "orderId": int(record.get("orderId") or 0),
                    "tradeId": None,
                    "price": float(record.get("price") or 0),
                    "qty": float(record.get("qty") or 0),
                    "quoteQty": float(record.get("quoteQty") or 0),
                    "commission": float(record.get("commission") or 0),
                    "commissionAsset": str(record.get("commissionAsset") or ""),
                    "time": int(record["time"]),
                    "timestamp_human": str(record.get("timestamp_human") or ""),
                    "isBuyer": bool(record.get("isBuyer")),
                }
            )
        return cached
    except Exception:
        return []


def fetch_portfolio() -> tuple[
    list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]
]:
    api = require_client()
    account = api.get_account(recvWindow=10_000)
    try:
        earn_response = api.get_simple_earn_flexible_product_position(size=100)
        earn_data = earn_response.get("rows", [])
    except Exception:
        earn_data = []
    earn_assets = {
        item.get("asset") for item in earn_data if float(item.get("totalAmount", 0)) > 0
    }
    raw: list[tuple[str, float, str]] = []
    for item in account.get("balances", []):
        quantity = float(item.get("free", 0)) + float(item.get("locked", 0))
        if quantity > 0:
            normalized = normalize_asset(item["asset"])
            if item["asset"].startswith("LD") and normalized in earn_assets:
                continue
            source = "Earn" if item["asset"].startswith("LD") else "Spot"
            raw.append((item["asset"], quantity, source))

    for item in earn_data:
        quantity = float(item.get("totalAmount", 0))
        if quantity > 0:
            raw.append((item["asset"], quantity, "Earn"))

    copy_spot_quantities: dict[str, float] = {}
    try:
        wallet_balances = signed_request(
            "/sapi/v1/asset/wallet/balance",
            {"quoteAsset": "USDT", "needBalanceDetail": "true"},
        )
        copy_wallet = next(
            (
                wallet
                for wallet in wallet_balances
                if wallet.get("walletName") == "Copy Trading"
            ),
            {},
        )
    except Exception:
        copy_wallet = {}
    for item in copy_wallet.get("assetBalances", []):
        quantity = sum(
            float(item.get(field) or 0)
            for field in ("free", "locked", "freeze", "withdrawing")
        )
        if quantity > 0:
            asset = normalize_asset(item["asset"])
            copy_spot_quantities[asset] = copy_spot_quantities.get(asset, 0) + quantity
            raw.append((asset, quantity, "Copy Spot"))

    grouped: dict[str, dict[str, Any]] = {}
    for asset, quantity, source in raw:
        normalized = normalize_asset(asset)
        record = grouped.setdefault(
            normalized,
            {"quantity": 0.0, "sources": set(), "source_quantities": {}},
        )
        record["quantity"] += quantity
        record["sources"].add(source)
        record["source_quantities"][source] = (
            record["source_quantities"].get(source, 0) + quantity
        )

    ticker_map = {
        ticker["symbol"]: float(ticker["price"])
        for ticker in api.get_all_tickers()
    }
    try:
        ticker_24h_map = {ticker["symbol"]: ticker for ticker in api.get_ticker()}
    except Exception:
        ticker_24h_map = {}

    exchange_symbols = {
        item["symbol"]: item
        for item in api.get_exchange_info().get("symbols", [])
    }
    known_trade_symbols = read_known_trade_symbols()
    significant_assets: set[str] = set()
    for asset, holding in grouped.items():
        lookup_asset = normalize_for_lookup(asset)
        if asset in USD_LIKE_QUOTES:
            estimated_value = holding["quantity"]
        elif asset == "BRL":
            estimated_value = (
                holding["quantity"] / ticker_map["USDTBRL"]
                if ticker_map.get("USDTBRL")
                else 0
            )
        else:
            estimated_value = holding["quantity"] * ticker_map.get(
                f"{lookup_asset}USDT", 0
            )
        if estimated_value >= 1:
            significant_assets.add(lookup_asset)

    candidate_symbols = {
        symbol for symbol in known_trade_symbols if symbol in exchange_symbols
    }
    for symbol, metadata in exchange_symbols.items():
        base_asset = metadata.get("baseAsset")
        quote_asset = metadata.get("quoteAsset")
        if metadata.get("status") != "TRADING":
            continue
        # Query common cash quotes only for meaningful current positions.
        # Previously seen pairs remain included above, even after a full sale.
        if base_asset in significant_assets and quote_asset in TRADE_QUOTES:
            candidate_symbols.add(symbol)

    def symbol_trades(symbol: str) -> tuple[str, list[dict[str, Any]]]:
        try:
            return symbol, api.get_my_trades(
                symbol=symbol, limit=1000, recvWindow=10_000
            )
        except Exception:
            return symbol, []

    trades_by_symbol: dict[str, list[dict[str, Any]]] = {}
    with ThreadPoolExecutor(max_workers=10) as executor:
        for symbol, trades in executor.map(symbol_trades, sorted(candidate_symbols)):
            trades_by_symbol[symbol] = trades

    grouped_asset_by_lookup = {
        normalize_for_lookup(asset): asset for asset in grouped
    }
    trade_map: dict[str, list[dict[str, Any]]] = {
        asset: [] for asset in grouped
    }
    for symbol, trades in trades_by_symbol.items():
        metadata = exchange_symbols.get(symbol, {})
        base_asset = metadata.get("baseAsset") or asset_from_symbol(symbol)
        quote_asset = metadata.get("quoteAsset") or next(
            (quote for quote in TRADE_QUOTES if symbol.endswith(quote)), "USDT"
        )
        grouped_asset = grouped_asset_by_lookup.get(base_asset, base_asset)
        if quote_asset in USD_LIKE_QUOTES:
            trade_map.setdefault(grouped_asset, []).extend(trades)

    result: list[dict[str, Any]] = []
    for asset, holding in grouped.items():
        total_balance = holding["quantity"]
        if asset in {"USDT", "FDUSD", "USDC"}:
            current_price = 1.0
        elif asset == "BRL":
            current_price = 1 / ticker_map["USDTBRL"] if ticker_map.get("USDTBRL") else None
        else:
            current_price = ticker_map.get(f"{normalize_for_lookup(asset)}USDT")
        market_symbol = (
            "USDTBRL" if asset == "BRL" else f"{normalize_for_lookup(asset)}USDT"
        )
        market_24h = ticker_24h_map.get(market_symbol, {})
        price_change_percent_24h = 0.0 if asset in {"USDT", "FDUSD", "USDC"} else None
        if market_24h:
            raw_change = float(market_24h.get("priceChangePercent", 0))
            price_change_percent_24h = -raw_change if asset == "BRL" else raw_change
        current_value = total_balance * current_price if current_price is not None else None
        previous_value = (
            current_value / (1 + price_change_percent_24h / 100)
            if current_value is not None
            and price_change_percent_24h is not None
            and price_change_percent_24h != -100
            else current_value
        )
        value_change_24h = (
            current_value - previous_value
            if current_value is not None and previous_value is not None
            else None
        )
        average_buy_price = None
        total_cost = None
        copy_spot_quantity = holding["source_quantities"].get("Copy Spot", 0)
        known_quantity = total_balance - copy_spot_quantity
        cost_basis_unavailable = (
            copy_spot_quantity > 0
            and asset not in {"BRL", "USDT", "FDUSD", "USDC"}
        )
        if asset not in {"BRL", "USDT", "FDUSD", "USDC"}:
            buys = [trade for trade in trade_map.get(asset, []) if trade.get("isBuyer")]
            buy_qty = sum(float(trade["qty"]) for trade in buys)
            buy_value = sum(float(trade["quoteQty"]) for trade in buys)
            if buy_qty > 0:
                average_buy_price = buy_value / buy_qty
                total_cost = (
                    known_quantity if cost_basis_unavailable else total_balance
                ) * average_buy_price

        if asset in {"USDT", "FDUSD", "USDC"}:
            average_buy_price = 1.0
            total_cost = total_balance
        acquired_free = not total_cost and not cost_basis_unavailable
        known_current_value = (
            known_quantity * current_price
            if current_price is not None and cost_basis_unavailable
            else current_value
        )
        profit_loss = (
            None
            if cost_basis_unavailable and not total_cost
            else (
                known_current_value if acquired_free else known_current_value - total_cost
            ) if known_current_value is not None else None
        )
        profit_loss_percent = (
            None
            if cost_basis_unavailable and not total_cost
            else (
            100.0 if acquired_free and current_value is not None
            else (profit_loss / total_cost * 100 if total_cost else None)
            )
        )
        result.append(
            {
                "asset": asset,
                "total_balance": total_balance,
                "current_price_usdt": current_price,
                "current_value_usdt": current_value,
                "average_buy_price_usdt": average_buy_price,
                "total_cost_usdt": total_cost,
                "profit_loss_usdt": profit_loss,
                "profit_loss_percent": profit_loss_percent,
                "price_change_percent_24h": price_change_percent_24h,
                "value_change_24h_usdt": value_change_24h,
                "acquired_free": acquired_free,
                "cost_basis_unavailable": cost_basis_unavailable,
                "sources": sorted(holding["sources"]),
            }
        )
    balances = sorted(result, key=lambda item: item["current_value_usdt"] or 0, reverse=True)
    balance_by_asset = {item["asset"]: item for item in balances}
    copy_spot_balances = sorted(
        (
            {
                "asset": asset,
                "quantity": quantity,
                "current_price_usdt": balance_by_asset.get(asset, {}).get("current_price_usdt"),
                "current_value_usdt": (
                    quantity * balance_by_asset[asset]["current_price_usdt"]
                    if balance_by_asset.get(asset, {}).get("current_price_usdt") is not None
                    else None
                ),
            }
            for asset, quantity in copy_spot_quantities.items()
        ),
        key=lambda item: item["current_value_usdt"] or 0,
        reverse=True,
    )
    all_trades = read_cached_trades_from_workbook()
    seen_trades = {
        (
            item["symbol"],
            item.get("orderId"),
            item["time"],
            item["qty"],
            item["price"],
        )
        for item in all_trades
    }
    for symbol, trades in trades_by_symbol.items():
        metadata = exchange_symbols.get(symbol, {})
        base_asset = metadata.get("baseAsset") or asset_from_symbol(symbol)
        quote_asset = metadata.get("quoteAsset") or next(
            (quote for quote in TRADE_QUOTES if symbol.endswith(quote)), "USDT"
        )
        asset = grouped_asset_by_lookup.get(base_asset, base_asset)
        for trade in trades:
            trade_id = int(trade["id"]) if trade.get("id") is not None else None
            identity = (
                symbol,
                int(trade.get("orderId", 0)),
                int(trade["time"]),
                float(trade["qty"]),
                float(trade["price"]),
            )
            if identity in seen_trades:
                continue
            seen_trades.add(identity)
            timestamp_ms = int(trade["time"])
            all_trades.append(
                {
                    "asset": asset,
                    "symbol": symbol,
                    "quoteAsset": quote_asset,
                    "orderId": trade["orderId"],
                    "tradeId": trade_id,
                    "price": float(trade["price"]),
                    "qty": float(trade["qty"]),
                    "quoteQty": float(trade["quoteQty"]),
                    "commission": float(trade["commission"]),
                    "commissionAsset": trade["commissionAsset"],
                    "time": timestamp_ms,
                    "timestamp_human": datetime.fromtimestamp(
                        timestamp_ms / 1000, tz=timezone.utc
                    ).strftime("%Y-%m-%d %H:%M:%S"),
                    "isBuyer": bool(trade["isBuyer"]),
                }
            )
    transactions = sorted(all_trades, key=lambda item: item["time"], reverse=True)
    return balances, transactions, copy_spot_balances


def asset_from_symbol(symbol: str) -> str:
    for quote in ("USDT", "FDUSD", "USDC", "BTC", "ETH", "BNB", "BRL"):
        if symbol.endswith(quote) and len(symbol) > len(quote):
            return symbol[: -len(quote)]
    return symbol


def fetch_open_orders() -> list[dict[str, Any]]:
    try:
        orders = require_client().get_open_orders(recvWindow=10_000)
    except Exception:
        return []
    result = []
    for item in orders:
        timestamp_ms = int(item.get("time") or item.get("updateTime") or 0)
        result.append(
            {
                "symbol": item.get("symbol", ""),
                "asset": asset_from_symbol(item.get("symbol", "")),
                "order_id": item.get("orderId"),
                "side": item.get("side"),
                "type": item.get("type"),
                "price": float(item.get("price") or 0),
                "quantity": float(item.get("origQty") or 0),
                "executed_quantity": float(item.get("executedQty") or 0),
                "status": item.get("status"),
                "time": timestamp_ms,
            }
        )
    return sorted(result, key=lambda item: item["time"], reverse=True)


def parse_binance_time(value: Any) -> tuple[int, str]:
    if value in (None, ""):
        return 0, "—"
    if isinstance(value, (int, float)) or str(value).isdigit():
        timestamp_ms = int(value)
        return timestamp_ms, datetime.fromtimestamp(
            timestamp_ms / 1000, tz=timezone.utc
        ).strftime("%Y-%m-%d %H:%M:%S")
    raw = str(value).replace("T", " ").replace("Z", "")
    try:
        parsed = datetime.fromisoformat(raw).replace(tzinfo=timezone.utc)
        return int(parsed.timestamp() * 1000), parsed.strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return 0, raw


def fetch_crypto_movements() -> list[dict[str, Any]]:
    api = require_client()
    movements: list[dict[str, Any]] = []
    try:
        deposits = api.get_deposit_history(limit=1000, recvWindow=10_000)
    except Exception:
        deposits = []
    for item in deposits:
        timestamp_ms, timestamp_human = parse_binance_time(
            item.get("completeTime") or item.get("insertTime")
        )
        movements.append(
            {
                "type": "deposit",
                "asset": item.get("coin", ""),
                "amount": float(item.get("amount") or 0),
                "network": item.get("network") or "—",
                "status": "Concluído" if item.get("status") == 1 else "Pendente",
                "tx_id": item.get("txId") or "",
                "time": timestamp_ms,
                "timestamp_human": timestamp_human,
            }
        )
    withdraw_status = {
        0: "E-mail enviado", 1: "Cancelado", 2: "Aguardando aprovação",
        3: "Rejeitado", 4: "Processando", 5: "Falha", 6: "Concluído",
    }
    try:
        withdrawals = api.get_withdraw_history(limit=1000, recvWindow=10_000)
    except Exception:
        withdrawals = []
    for item in withdrawals:
        timestamp_ms, timestamp_human = parse_binance_time(
            item.get("completeTime") or item.get("applyTime")
        )
        movements.append(
            {
                "type": "withdrawal",
                "asset": item.get("coin", ""),
                "amount": float(item.get("amount") or 0),
                "network": item.get("network") or "—",
                "status": withdraw_status.get(item.get("status"), str(item.get("status", "—"))),
                "tx_id": item.get("txId") or "",
                "time": timestamp_ms,
                "timestamp_human": timestamp_human,
            }
        )
    return sorted(movements, key=lambda item: item["time"], reverse=True)


def build_insights(balances: list[dict[str, Any]]) -> dict[str, Any]:
    valued = [item for item in balances if (item.get("current_value_usdt") or 0) > 0]
    total = sum(item["current_value_usdt"] for item in valued)
    weights = [item["current_value_usdt"] / total for item in valued] if total else []
    stablecoins = {"USDT", "USDC", "FDUSD", "DAI", "TUSD"}
    change_24h = sum(item.get("value_change_24h_usdt") or 0 for item in valued)
    previous_total = total - change_24h
    return {
        "change_24h_usdt": change_24h,
        "change_24h_percent": change_24h / previous_total * 100 if previous_total else 0,
        "largest_asset": valued[0]["asset"] if valued else None,
        "largest_position_percent": weights[0] * 100 if weights else 0,
        "top3_concentration_percent": sum(weights[:3]) * 100,
        "stablecoin_percent": sum(
            item["current_value_usdt"] for item in valued if item["asset"] in stablecoins
        ) / total * 100 if total else 0,
        "hhi": sum(weight * weight for weight in weights) * 10_000,
    }


def read_fiat_history_from_workbook() -> list[dict[str, Any]]:
    if not WORKBOOK_FILE.exists():
        return []
    workbook = load_workbook(WORKBOOK_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook["fiat_history"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").removeprefix("Column1.") for value in next(rows)]
        result = []
        for values in rows:
            raw = dict(zip(headers, values))
            timestamp = str(raw.get("timestamp_human") or "")
            if not timestamp:
                continue
            result.append(
                {
                    "type": raw.get("type"),
                    "orderId": raw.get("orderId"),
                    "fiatCurrency": raw.get("fiatCurrency"),
                    "indicatedAmount": raw.get("indicatedAmount"),
                    "amount": raw.get("amount"),
                    "totalFee": raw.get("totalFee"),
                    "status": raw.get("status"),
                    "method": raw.get("method"),
                    "createTime": 0,
                    "timestamp_human": timestamp,
                }
            )
        return result
    finally:
        workbook.close()


def fetch_fiat_history() -> list[dict[str, Any]]:
    records = read_fiat_history_from_workbook()
    window_ms = 89 * 24 * 60 * 60 * 1000
    now_ms = int(time.time() * 1000)
    cursor = now_ms - window_ms
    for path, record_type in (
        ("/sapi/v1/fiat/orders", "Order"),
        ("/sapi/v1/fiat/payments", "Deposit"),
    ):
        try:
            response = signed_request(
                path,
                {"transactionType": 0, "rows": 100, "page": 1, "beginTime": cursor, "endTime": now_ms},
            )
        except Exception:
            continue
        for item in response.get("data", []):
            timestamp_ms = int(item.get("createTime", 0))
            records.append(
                {
                    "type": record_type,
                    "orderId": item.get("orderId"),
                    "fiatCurrency": item.get("fiatCurrency"),
                    "indicatedAmount": item.get("indicatedAmount"),
                    "amount": item.get("amount"),
                    "totalFee": item.get("totalFee"),
                    "status": item.get("status"),
                    "method": item.get("paymentMethod"),
                    "createTime": timestamp_ms,
                    "timestamp_human": datetime.fromtimestamp(
                        timestamp_ms / 1000, tz=timezone.utc
                    ).strftime("%Y-%m-%d %H:%M:%S"),
                }
            )
    deduplicated = {
        (item.get("type"), item.get("orderId"), item.get("timestamp_human"), str(item.get("amount"))): item
        for item in records
    }
    return sorted(deduplicated.values(), key=lambda item: item["timestamp_human"], reverse=True)


def read_history() -> list[dict[str, Any]]:
    if not HISTORY_FILE.exists():
        return []
    rows: list[dict[str, Any]] = []
    with HISTORY_FILE.open("r", encoding="utf-8-sig", newline="") as file:
        for row in csv.DictReader(file):
            raw = row.get("datetime", "")
            parsed = None
            for pattern in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M"):
                try:
                    parsed = datetime.strptime(raw, pattern)
                    break
                except ValueError:
                    continue
            if parsed is None:
                continue
            try:
                value = float(row.get("total_portfolio_value_usdt", 0))
            except (TypeError, ValueError):
                continue
            rows.append({"datetime": parsed.isoformat(), "total_portfolio_value_usdt": value})
    return sorted(rows, key=lambda item: item["datetime"])


def write_history(total_value: float) -> None:
    now = datetime.now()
    existing = read_history()
    today = now.date().isoformat()
    new_row = {"datetime": now.replace(microsecond=0).isoformat(), "total_portfolio_value_usdt": total_value}
    filtered = [row for row in existing if not row["datetime"].startswith(today)]
    filtered.append(new_row)
    with HISTORY_FILE.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["datetime", "total_portfolio_value_usdt"])
        writer.writeheader()
        for row in sorted(filtered, key=lambda item: item["datetime"]):
            writer.writerow(row)


def refresh_excel_workbook() -> None:
    if not WORKBOOK_FILE.exists() or not EXCEL_REFRESH_SCRIPT.exists():
        return
    try:
        subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(EXCEL_REFRESH_SCRIPT), "-WorkbookPath", str(WORKBOOK_FILE)],
            cwd=BASE_DIR,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
    except Exception:
        pass


def refresh_excel_in_background() -> None:
    threading.Thread(target=refresh_excel_workbook, daemon=True).start()


def build_dashboard(force: bool = False) -> dict[str, Any]:
    with _cache_lock:
        cached = _dashboard_cache["data"]
        fresh = time.monotonic() - _dashboard_cache["updated_monotonic"] < CACHE_SECONDS
        if cached is not None and fresh and not force:
            return cached
    if not _refresh_lock.acquire(blocking=False):
        with _refresh_lock:
            with _cache_lock:
                if _dashboard_cache["data"] is not None:
                    return _dashboard_cache["data"]
        raise RuntimeError("A atualização dos dados não foi concluída.")
    try:
        with ThreadPoolExecutor(max_workers=4) as executor:
            portfolio_future = executor.submit(fetch_portfolio)
            fiat_future = executor.submit(fetch_fiat_history)
            orders_future = executor.submit(fetch_open_orders)
            movements_future = executor.submit(fetch_crypto_movements)
            balances, transactions, copy_spot_balances = portfolio_future.result()
            try:
                fiat_history = fiat_future.result()
            except Exception:
                fiat_history = read_fiat_history_from_workbook()
            open_orders = orders_future.result()
            crypto_movements = movements_future.result()
        portfolio_value = sum(item["current_value_usdt"] or 0 for item in balances)
        invested_cost = sum(item["total_cost_usdt"] or 0 for item in balances)
        pnl = sum(item["profit_loss_usdt"] or 0 for item in balances)
        roi = pnl / invested_cost * 100 if invested_cost else 0
        usdt_brl = get_current_price("USDTBRL")
        write_history(portfolio_value)
        payload = {
            "updated_at": datetime.now().astimezone().isoformat(),
            "portfolio_value_usdt": portfolio_value,
            "portfolio_value_brl": portfolio_value * usdt_brl if usdt_brl else None,
            "usdt_brl": usdt_brl,
            "invested_cost_usdt": invested_cost,
            "profit_loss_usdt": pnl,
            "roi_percent": roi,
            "balances": balances,
            "copy_spot_balances": copy_spot_balances,
            "transactions": transactions,
            "fiat_history": fiat_history,
            "open_orders": open_orders,
            "crypto_movements": crypto_movements,
            "insights": build_insights(balances),
            "history": read_history(),
        }
        with _cache_lock:
            _dashboard_cache["data"] = payload
            _dashboard_cache["updated_monotonic"] = time.monotonic()
        return payload
    finally:
        _refresh_lock.release()


@app.get("/", include_in_schema=False)
def dashboard_page() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "app": "crypto-investments-dashboard", "credentials_configured": client is not None}


@app.get("/api/dashboard")
def dashboard_data(refresh: bool = Query(default=False)) -> dict[str, Any]:
    try:
        data = build_dashboard(force=refresh)
        if refresh:
            refresh_excel_in_background()
        return data
    except Exception as exc:
        with _cache_lock:
            stale = _dashboard_cache["data"]
        if stale is not None:
            return {**stale, "warning": str(exc), "stale": True}
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.post("/api/refresh")
def refresh_data() -> dict[str, Any]:
    try:
        data = build_dashboard(force=True)
        refresh_excel_in_background()
        return data
    except Exception as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.get("/balances")
def balances_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["balances"]


@app.get("/transactions")
def transactions_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["transactions"]


@app.get("/fiat_history")
def fiat_history_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["fiat_history"]


@app.get("/api/market-24h")
def market_24h_endpoint() -> list[dict[str, Any]]:
    return [
        {
            "asset": item["asset"],
            "price_change_percent_24h": item["price_change_percent_24h"],
            "value_change_24h_usdt": item["value_change_24h_usdt"],
        }
        for item in build_dashboard()["balances"]
    ]


@app.get("/api/open-orders")
def open_orders_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["open_orders"]


@app.get("/api/crypto-movements")
def crypto_movements_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["crypto_movements"]


@app.get("/api/copy-spot")
def copy_spot_endpoint() -> list[dict[str, Any]]:
    return build_dashboard()["copy_spot_balances"]
